import openpyxl
import tempfile
import os
import pandas as pd

def generate_template_excel(df: pd.DataFrame, template_path: str) -> str:
    """
    Loads an existing template.xlsx, injects the DataFrame into the 'RAW_DATA' sheet,
    and returns the path to the newly generated temporary .xlsx file.
    """
    try:
        # Load the user's template
        wb = openpyxl.load_workbook(template_path)
    except FileNotFoundError:
        # Fallback if user hasn't created the template yet, we create a basic one
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "RAW_DATA"
    except Exception as e:
        raise ValueError(f"Failed to open template.xlsx. Ensure it is a valid Excel file. Error: {str(e)}")

    # Ensure RAW_DATA sheet exists
    sheet_name = "RAW_DATA"
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        # Clear existing data except header (row 1)
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
    else:
        ws = wb.create_sheet(sheet_name)

    # Write Headers
    headers = df.columns.tolist()
    for col_num, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_num, value=str(header))

    # Write Data
    # Iterating over values is much faster than iterrows()
    for row_num, row_data in enumerate(df.values, 2):
        for col_num, value in enumerate(row_data, 1):
            # openpyxl handles native python types well, but ensure no complex types
            ws.cell(row=row_num, column=col_num, value=value)

    # Save to a temporary file
    fd, temp_path = tempfile.mkstemp(suffix=".xlsx", prefix="nexus_report_")
    os.close(fd)
    
    wb.save(temp_path)
    return temp_path
